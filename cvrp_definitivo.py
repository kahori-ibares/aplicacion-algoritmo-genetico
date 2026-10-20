import random   
import math     
import copy     
import matplotlib          
matplotlib.use("Agg")      
import matplotlib.pyplot as plt          
import matplotlib.patches as mpatches   
from openpyxl import Workbook                              
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter  
from openpyxl.chart import BarChart, Reference 

DEPOSITO = {"nombre": "Paso Texca", "lat": 16.98188, "lon": -99.81288}
NUM_CAM    = 12     
CAP        = 10.0   
MAX_VIAJES = 3      
TOTAL_BAS  = 250.0  
k = 4
tam = 150
gens = 500
elite = 6
seed = 42
COLONIAS = [
    "Centro",           "Renacimiento",         "Zapata",
    "Coloso",           "Costa Azul",            "Icacos",
    "Progreso",         "Mozimba",               "Caleta",
    "Caletilla",        "Pie de la Cuesta",      "Garita",
    "Hornos",           "La Sabana",             "Llano Largo",
    "Las Cruces",       "Emiliano Zapata",       "Jardin Azteca",
    "Jardin Mangos",    "Jardin Palmas",         "Cumbres de Figueroa",
    "Farallon",         "La Mira",               "Las Playas",
    "Puerto Marques",   "El Coloso II",          "Ciudad San Agustin",
    "Pedregoso",        "Lomas de Chapultepec",  "Bella Vista",
    "Vista Alegre",     "Morelos",               "Cuauhtemoc",
    "La Laja",          "Alta Progreso",         "Santa Cruz",
    "Libertadores"
]
NUM_COL = len(COLONIAS)
COORDS = [
    (16.8634,-99.8825),    (16.8780,-99.9020),    (16.8590,-99.8920),  
    (16.8700,-99.9150),    (16.8330,-99.8640),    (16.8240,-99.8540),  
    (16.8820,-99.9100),    (16.8750,-99.9200),    (16.8430,-99.9080),  
    (16.8410,-99.9100),    (16.9100,-99.9400),    (16.8650,-99.9050),  
    (16.8560,-99.8850),    (16.9050,-99.8960),    (16.9300,-99.8800),
    (16.8850,-99.8750),    (16.8700,-99.8950),    (16.8900,-99.9100),  
    (16.8800,-99.9050),    (16.8760,-99.8980),    (16.8650,-99.9180), 
    (16.8500,-99.8900),    (16.8480,-99.9020),    (16.8320,-99.8720),  
    (16.7980,-99.8840),    (16.8720,-99.9130),    (16.8850,-99.9050), 
    (16.8930,-99.8900),    (16.8550,-99.9200),    (16.8610,-99.8960),  
    (16.8580,-99.8870),    (16.8710,-99.8880),    (16.8670,-99.8810),  
    (16.9200,-99.8700),    (16.8960,-99.8850),    (16.8440,-99.9140),  
    (16.8800,-99.9300),  
]

random.seed(42)
_r = [random.uniform(3.0, 7.0) for _ in range(NUM_COL)]
_f = TOTAL_BAS / sum(_r) 
DEMANDA = [round(x * _f, 2) for x in _r] 
DEMANDA[-1] = round(DEMANDA[-1] + round(TOTAL_BAS - sum(DEMANDA), 2), 2)
assert max(DEMANDA) <= CAP + 1e-6, f"Colonia con demanda > {CAP}t no permitida"

def hav(la1, lo1, la2, lo2):
    R = 6371.0     
    p1 = math.radians(la1)
    p2 = math.radians(la2)
    a = (math.sin(math.radians(la2 - la1) / 2) ** 2 +
         math.cos(p1) * math.cos(p2) *
         math.sin(math.radians(lo2 - lo1) / 2) ** 2)
    return R * 2 * math.asin(math.sqrt(max(0, a)))

NODOS = [(DEPOSITO["lat"], DEPOSITO["lon"])] + list(COORDS)
N = len(NODOS)   

DIST = [
    [hav(NODOS[i][0], NODOS[i][1], NODOS[j][0], NODOS[j][1])
     if i != j else 0.0
     for j in range(N)]
    for i in range(N)
]


def decodificar(cromosoma):
    camiones = [[{"nodos": [], "carga": 0.0}] for _ in range(NUM_CAM)]
    cam_ptr   = 0   # Índice del camión que está siendo llenado actualmente
    viaje_ptr = 0   # Índice del viaje dentro del camión actual

    for nodo in cromosoma:
        dem = DEMANDA[nodo - 1]
        colocado = False   
        for _ in range(NUM_CAM * MAX_VIAJES):
            viaje = camiones[cam_ptr][viaje_ptr]
            if viaje["carga"] + dem <= CAP + 1e-9:
                viaje["nodos"].append(nodo)
                viaje["carga"] = round(viaje["carga"] + dem, 4)
                colocado = True
                break   

            else:
                if viaje_ptr + 1 < MAX_VIAJES:
                    viaje_ptr += 1
                    camiones[cam_ptr].append({"nodos": [], "carga": 0.0})

                else:
                    cam_ptr  += 1   
                    viaje_ptr = 0   
                    if cam_ptr >= NUM_CAM:
                        best_c, best_v, best_esp = 0, 0, -1
                        for c in range(NUM_CAM):
                            for v in range(len(camiones[c])):
                                esp = CAP - camiones[c][v]["carga"]
                                if esp > best_esp:
                                    best_c, best_v, best_esp = c, v, esp
                        camiones[best_c][best_v]["nodos"].append(nodo)
                        camiones[best_c][best_v]["carga"] = round(
                            camiones[best_c][best_v]["carga"] + dem, 4)
                        colocado = True
                        break
                    if not camiones[cam_ptr]:
                        camiones[cam_ptr] = [{"nodos": [], "carga": 0.0}]
        if not colocado:
            pass
    for cam in camiones:
        to_del = [v for v in cam if not v["nodos"]]
        for v in to_del:
            cam.remove(v)
    return camiones   

def dist_viaje(nodos):
    if not nodos:
        return 0.0  
    d = DIST[0][nodos[0]]
    for i in range(len(nodos) - 1):
        d += DIST[nodos[i]][nodos[i + 1]]
    d += DIST[nodos[-1]][0]
    return d   

def fitness(crom):
    total = sum(
        dist_viaje(v["nodos"])
        for cam in decodificar(crom)   
        for v in cam                   
    )
    return 1.0 / (total + 1e-9)

def nuevo_ind():
    x = list(range(1, NUM_COL + 1))   
    random.shuffle(x)                  
    return x                           
def torneo(pop, fits, k):
    cs = random.sample(range(len(pop)), k)
    return copy.copy(pop[max(cs, key=lambda i: fits[i])])

def ox(p1, p2):
    n = len(p1)
    a, b = sorted(random.sample(range(n), 2))
    h = [None] * n
    h[a:b + 1] = p1[a:b + 1]
    ptr = (b + 1) % n
    for g in p2[b + 1:] + p2[:b + 1]:   
        if g not in h:                    
            h[ptr] = g
            ptr = (ptr + 1) % n
    return h

def mut_swap(ind, t=0.03):
    ind = copy.copy(ind)           
    for i in range(len(ind)):
        if random.random() < t:    
            j = random.randint(0, len(ind) - 1)
            ind[i], ind[j] = ind[j], ind[i]   
    return ind

def mut_2opt(ind, t=0.07):
    if random.random() > t:
        return ind                  
    ind = copy.copy(ind)
    a, b = sorted(random.sample(range(len(ind)), 2))
    ind[a:b + 1] = reversed(ind[a:b + 1])   
    return ind

def ga(tam, gens, elite, seed):
    random.seed(seed)   
    pop = [nuevo_ind() for _ in range(tam)]
    print("  Gen    Dist. total (km)")
    print("  -----  ----------------")
    for g in range(gens):   
        fits = [fitness(i) for i in pop]
        orden = sorted(range(tam), key=lambda i: fits[i], reverse=True)
        nueva_gens = [copy.copy(pop[i]) for i in orden[:elite]]
        while len(nueva_gens) < tam:
            p1 = torneo(pop, fits, k)   
            p2 = torneo(pop, fits, k)   
            h = ox(p1, p2) if random.random() < 0.85 else copy.copy(p1)
            h = mut_swap(h)   
            h = mut_2opt(h)   
            nueva_gens.append(h)  
        pop = nueva_gens   
        if (g + 1) % 100 == 0:
            print(f"  {g+1:>5}  {1.0/max(fits):>16,.2f}")
    fits = [fitness(i) for i in pop]
    return pop[max(range(tam), key=lambda i: fits[i])]

PALETTE = [
    "#E74C3C", "#3498DB", "#2ECC71", "#F39C12", "#9B59B6",
    "#1ABC9C", "#E67E22", "#95A5A6", "#F1C40F", "#16A085",
    "#8E44AD", "#27AE60", "#D35400", "#2980B9", "#C0392B",
    "#7F8C8D", "#52BE80", "#5DADE2", "#F0B27A", "#C39BD3",
    "#76D7C4", "#F7DC6F", "#FDEBD0", "#AED6F1", "#FADBD8",
]

def brd():
    """Retorna un objeto Border con líneas delgadas en los 4 lados."""
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def fill(hex6):
    """
    Convierte un color hex CSS (#RRGGBB o RRGGBB) al formato ARGB que
    necesita openpyxl (FFRRGGBB, donde FF = opacidad máxima).
    """
    return PatternFill("solid", fgColor="FF" + hex6.lstrip("#").upper())

def hfont(sz=10):
    """Fuente de encabezado: Arial, blanca, negrita."""
    return Font(bold=True, color="FFFFFF", name="Arial", size=sz)

def bfont(sz=9, bold=False, color="000000"):
    """Fuente de cuerpo: Arial, negra, tamaño configurable."""
    return Font(bold=bold, name="Arial", size=sz, color=color)


# ──────────────────────────────────────────────────────────────
#  SECCIÓN 8 — GENERACIÓN DEL EXCEL DETALLADO
# ──────────────────────────────────────────────────────────────

def make_excel(camiones, path):
    """
    Genera el archivo Excel con 3 hojas:
      1) Rutas Detalladas: cada fila = un viaje con ruta completa, toneladas y km
      2) Dashboard: KPIs por camión + gráfica de barras
      3) Demanda Colonias: tabla de referencia con toneladas por colonia
    """
    wb = Workbook()   # Crear libro Excel nuevo (vacío)

    # ── HOJA 1: RUTAS DETALLADAS ──────────────────────────────────
    ws = wb.active          # Primera hoja (creada automáticamente)
    ws.title = "Rutas Detalladas"

    # Colores del diseño (hexadecimales sin #)
    DARK = "1A1A2E"   # Azul marino oscuro — encabezados principales
    MED  = "16213E"   # Azul marino medio — encabezados de columna

    # Fila 1: Título que ocupa todas las columnas (A a I)
    ws.merge_cells("A1:I1")   # Une las celdas A1 hasta I1 en una sola
    ws["A1"] = "🚛  RUTEO DE CAMIONES DE BASURA — PUERTO DE ACAPULCO, GRO."
    ws["A1"].font = Font(bold=True, name="Arial", size=15, color="FFFFFF")
    ws["A1"].fill = fill(DARK)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36   # Altura de la fila en puntos

    # Fila 2: Información general del problema
    ws.merge_cells("A2:I2")
    ws["A2"] = (f"Depósito: {DEPOSITO['nombre']}   |   "
                f"Camiones: {NUM_CAM}   |   Capacidad: {CAP} t/viaje   |   "
                f"Máx. viajes: {MAX_VIAJES}   |   Basura total: {TOTAL_BAS} t")
    ws["A2"].font = Font(italic=True, name="Arial", size=9, color="AAAAAA")
    ws["A2"].fill = fill(DARK)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 8   # Fila 3: espacio visual vacío

    # Fila 4: Encabezados de las columnas de datos
    hdrs = ["Camión", "Color", "# Viaje",
            "Salida", "Colonias visitadas (ruta completa)", "Llegada",
            "Ton. este viaje", "Ton. acumuladas", "Km est. viaje"]
    col_w = [11, 5, 8, 14, 55, 14, 14, 14, 12]   # Ancho en caracteres de cada columna

    for ci, h in enumerate(hdrs, 1):   # enumerate(hdrs,1) da índice 1-based (A=1)
        c = ws.cell(row=4, column=ci, value=h)
        c.fill = fill(MED)
        c.font = hfont(9)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = brd()
    ws.row_dimensions[4].height = 30

    # Aplicar ancho a cada columna usando get_column_letter (1→"A", 2→"B", etc.)
    for ci, w in enumerate(col_w, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # Variables acumuladoras para el resumen global
    row = 5          # La primera fila de datos empieza en la 5
    tots_ton = []    # Toneladas totales por camión
    tots_km  = []    # Kilómetros totales por camión
    tots_viaj = []   # Número de viajes por camión
    used_cams = 0    # Contador de camiones que hicieron al menos un viaje

    # ── Iterar sobre cada camión ──────────────────────────────
    for ci, cam in enumerate(camiones):
        # Filtrar viajes vacíos (camiones sin asignación)
        viajes = [v for v in cam if v["nodos"]]
        nv = len(viajes)   # Número de viajes reales del camión

        if nv == 0:
            # Camión sin rutas → agregar ceros al resumen y continuar
            tots_ton.append(0.0); tots_km.append(0.0); tots_viaj.append(0)
            continue

        used_cams += 1
        color_hex = PALETTE[ci % len(PALETTE)].lstrip("#")   # Color del camión (sin #)
        cam_fill  = fill(color_hex)    # Relleno de color para la columna "Color"
        ton_ac = 0.0; km_ac = 0.0     # Acumuladores para este camión

        # ── Iterar sobre cada viaje del camión ──────────────
        for vi, viaje in enumerate(viajes):
            # Construir la cadena de texto con la ruta del viaje
            # Ej: "Centro → Hornos → Farallon"
            nombres = COLONIAS[viaje["nodos"][0] - 1]
            if len(viaje["nodos"]) > 1:
                for nn in viaje["nodos"][1:]:
                    nombres += " → " + COLONIAS[nn - 1]

            ton = round(viaje["carga"], 2)             # Toneladas cargadas en este viaje
            km  = round(dist_viaje(viaje["nodos"]), 2) # km recorridos en este viaje
            ton_ac = round(ton_ac + ton, 2)            # Acumular toneladas
            km_ac  = round(km_ac + km, 2)              # Acumular kilómetros

            # Función local para escribir y formatear una celda de manera concisa.
            # Se define DENTRO del loop para capturar `row` del contexto.
            def sc(col, val, **kw):
                c = ws.cell(row=row, column=col, value=val)
                c.border = brd()
                c.font   = bfont(9)
                c.alignment = Alignment(
                    horizontal=kw.get("align", "center"),
                    wrap_text=kw.get("wrap", False))
                if kw.get("bold"):  c.font = bfont(9, bold=True)
                if kw.get("bg"):    c.fill = fill(kw["bg"])
                if kw.get("fcol"):  c.font = Font(bold=kw.get("bold", False),
                                                  name="Arial", size=9, color=kw["fcol"])
                return c

            # Escribir cada columna de la fila del viaje
            sc(1, f"Camión {ci+1:02d}", bold=True, fcol=color_hex)   # Nombre camión en su color
            ws.cell(row=row, column=2).fill   = cam_fill              # Bloque de color sólido
            ws.cell(row=row, column=2).border = brd()
            sc(3, f"Viaje {vi+1}")                                     # Número de viaje
            sc(4, DEPOSITO["nombre"], bg="1A1A2E", fcol="FFD700")     # Salida (depósito, dorado)
            sc(5, nombres, align="left", wrap=True)                    # Ruta de colonias
            sc(6, DEPOSITO["nombre"], bg="1A1A2E", fcol="FFD700")     # Llegada (depósito, dorado)
            sc(7, ton); ws.cell(row=row, column=7).number_format = '0.00" t"'
            sc(8, ton_ac); ws.cell(row=row, column=8).number_format = '0.00" t"'
            sc(9, km);  ws.cell(row=row, column=9).number_format = '0.00" km"'

            # Altura dinámica: a más colonias en la ruta, más alta la fila
            ws.row_dimensions[row].height = max(15, 13 * len(viaje["nodos"]) // 3 + 14)
            row += 1   # Avanzar a la siguiente fila del Excel

        # ── Fila de subtotal por camión ──────────────────────
        GOLD_FILL = fill("FFF3CD")   # Amarillo claro para destacar subtotales
        for c in range(1, 10):
            ws.cell(row=row, column=c).fill   = GOLD_FILL
            ws.cell(row=row, column=c).border = brd()

        ws.cell(row=row, column=1).value = f"Camión {ci+1:02d}"
        ws.cell(row=row, column=1).font  = bfont(9, bold=True, color=color_hex)
        ws.cell(row=row, column=2).fill  = cam_fill
        ws.cell(row=row, column=3).value = f"SUBTOTAL ({nv} viaje{'s' if nv > 1 else ''})"
        ws.cell(row=row, column=3).font  = bfont(9, bold=True)
        ws.cell(row=row, column=7).value = ton_ac
        ws.cell(row=row, column=7).number_format = '0.00" t"'
        ws.cell(row=row, column=7).font  = bfont(9, bold=True)
        ws.cell(row=row, column=9).value = km_ac
        ws.cell(row=row, column=9).number_format = '0.00" km"'
        ws.cell(row=row, column=9).font  = bfont(9, bold=True)

        # Guardar resultados del camión para el resumen final
        tots_ton.append(ton_ac)
        tots_km.append(km_ac)
        tots_viaj.append(nv)
        row += 2   # Dejar una fila en blanco entre camiones (separador visual)

    # ── Fila de gran total (toda la operación) ───────────────
    for c in range(1, 10):
        ws.cell(row=row, column=c).fill   = fill(DARK)
        ws.cell(row=row, column=c).border = brd()
    ws.cell(row=row, column=1).value = f"{sum(tots_viaj)} viajes"
    ws.cell(row=row, column=1).font  = hfont(9)
    ws.cell(row=row, column=3).value = "GRAN TOTAL"
    ws.cell(row=row, column=3).font  = hfont(10)
    ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=4).value = f"{used_cams} camiones activos"
    ws.cell(row=row, column=4).font  = hfont(9)
    ws.cell(row=row, column=7).value = round(sum(tots_ton), 2)
    ws.cell(row=row, column=7).font  = hfont(10)
    ws.cell(row=row, column=7).number_format = '0.00" t"'
    ws.cell(row=row, column=9).value = round(sum(tots_km), 2)
    ws.cell(row=row, column=9).font  = hfont(10)
    ws.cell(row=row, column=9).number_format = '0.00" km"'
    ws.row_dimensions[row].height = 24

    # ── HOJA 2: DASHBOARD ────────────────────────────────────
    ws2 = wb.create_sheet("Dashboard")
    DARK2 = "1A1A2E"

    ws2.merge_cells("A1:F1")
    ws2["A1"] = "Dashboard — Resumen por camión"
    ws2["A1"].font = Font(bold=True, name="Arial", size=13, color="FFFFFF")
    ws2["A1"].fill = fill(DARK2)
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30

    # Encabezados del dashboard
    h2 = ["Camión", "Viajes", "Ton. total", "Km total", "% Basura", "Estado"]
    for ci, h in enumerate(h2, 1):
        c = ws2.cell(row=2, column=ci, value=h)
        c.fill = fill("16213E"); c.font = hfont(9); c.border = brd()
        c.alignment = Alignment(horizontal="center")

    # Anchos de columna del dashboard
    ws2.column_dimensions["A"].width = 12; ws2.column_dimensions["B"].width = 9
    ws2.column_dimensions["C"].width = 13; ws2.column_dimensions["D"].width = 13
    ws2.column_dimensions["E"].width = 12; ws2.column_dimensions["F"].width = 14

    active_r = 3   # Primera fila de datos

    for ci in range(NUM_CAM):
        if tots_viaj[ci] == 0:
            continue   # Saltar camiones sin actividad

        chex = PALETTE[ci % len(PALETTE)].lstrip("#")
        alt  = "F0F4F8" if active_r % 2 == 0 else "FFFFFF"   # Filas alternadas

        cells = [ws2.cell(row=active_r, column=j) for j in range(1, 7)]
        for c in cells: c.fill = fill(alt); c.border = brd(); c.font = bfont(9)

        cells[0].value = f"Camión {ci+1:02d}"
        cells[0].font  = Font(bold=True, name="Arial", size=9, color=chex)
        cells[1].value = tots_viaj[ci]
        cells[1].alignment = Alignment(horizontal="center")
        cells[2].value = tots_ton[ci]; cells[2].number_format = '0.00" t"'
        cells[2].alignment = Alignment(horizontal="center")
        cells[3].value = tots_km[ci];  cells[3].number_format = '0.00" km"'
        cells[3].alignment = Alignment(horizontal="center")

        # Porcentaje del total de basura recolectada por este camión
        pct = round(tots_ton[ci] / TOTAL_BAS * 100, 1)
        cells[4].value  = pct / 100   # openpyxl espera fracción decimal para formato %
        cells[4].number_format = '0.0%'
        cells[4].alignment = Alignment(horizontal="center")

        # Emoji de estado según número de viajes
        estado = ("🟢 Completo" if tots_viaj[ci] == MAX_VIAJES
                  else ("🟡 Parcial" if tots_viaj[ci] > 0 else "⚪ Libre"))
        cells[5].value = estado
        cells[5].alignment = Alignment(horizontal="center")
        active_r += 1

    # Fila de totales del dashboard (fila de cierre azul oscuro)
    for col_idx, val, fmt in [
        (1, "TOTAL",               None),
        (2, sum(tots_viaj),        None),
        (3, round(sum(tots_ton),2),'0.00" t"'),
        (4, round(sum(tots_km),2), '0.00" km"'),
    ]:
        c = ws2.cell(row=active_r + 1, column=col_idx, value=val)
        c.font = hfont(9); c.fill = fill(DARK2); c.border = brd()
        c.alignment = Alignment(horizontal="center")
        if fmt: c.number_format = fmt

    # ── Gráfica de barras de toneladas ────────────────────────
    chart = BarChart()
    chart.type  = "col"                                  # Barras verticales
    chart.title = "Toneladas recolectadas por camión"
    chart.y_axis.title = "Toneladas"
    chart.x_axis.title = "Camión"
    chart.height = 14; chart.width = 24

    # Reference define qué celdas del Excel alimentan la gráfica
    data_ref = Reference(ws2, min_col=3, min_row=2, max_row=active_r - 1)
    cats_ref = Reference(ws2, min_col=1, min_row=3, max_row=active_r - 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws2.add_chart(chart, "H2")   # Insertar la gráfica en la celda H2

    # ── HOJA 3: DEMANDA POR COLONIA ──────────────────────────
    ws3 = wb.create_sheet("Demanda Colonias")
    ws3.merge_cells("A1:C1")
    ws3["A1"] = "Demanda de basura por colonia"
    ws3["A1"].font = Font(bold=True, name="Arial", size=11, color="FFFFFF")
    ws3["A1"].fill = fill(DARK2)
    ws3["A1"].alignment = Alignment(horizontal="center")
    ws3.row_dimensions[1].height = 24

    for ci2, h in enumerate(["#", "Colonia", "Toneladas (t)"], 1):
        c = ws3.cell(row=2, column=ci2, value=h)
        c.fill = fill("16213E"); c.font = hfont(9); c.border = brd()
        c.alignment = Alignment(horizontal="center")

    ws3.column_dimensions["A"].width = 5
    ws3.column_dimensions["B"].width = 28
    ws3.column_dimensions["C"].width = 16

    for i, (col, dem) in enumerate(zip(COLONIAS, DEMANDA), 3):
        alt = "F8FAFB" if i % 2 == 0 else "FFFFFF"
        ws3.cell(row=i, column=1, value=i - 2).fill = fill(alt)
        ws3.cell(row=i, column=1).border = brd()
        ws3.cell(row=i, column=1).alignment = Alignment(horizontal="center")
        ws3.cell(row=i, column=1).font = bfont(9)
        ws3.cell(row=i, column=2, value=col).fill = fill(alt)
        ws3.cell(row=i, column=2).border = brd()
        ws3.cell(row=i, column=2).font = bfont(9)
        ws3.cell(row=i, column=3, value=dem).fill = fill(alt)
        ws3.cell(row=i, column=3).border = brd()
        ws3.cell(row=i, column=3).number_format = '0.00'
        ws3.cell(row=i, column=3).alignment = Alignment(horizontal="center")
        ws3.cell(row=i, column=3).font = bfont(9)

    lr = NUM_COL + 3   # Fila del total (después de todas las colonias)
    ws3.cell(row=lr, column=2, value="TOTAL").font = bfont(9, bold=True)
    ws3.cell(row=lr, column=2).fill = fill("FFF3CD"); ws3.cell(row=lr, column=2).border = brd()
    # Fórmula Excel para la suma (=SUM(C3:C39))
    ws3.cell(row=lr, column=3, value=f"=SUM(C3:C{NUM_COL+2})").font = bfont(9, bold=True)
    ws3.cell(row=lr, column=3).fill = fill("FFF3CD"); ws3.cell(row=lr, column=3).border = brd()
    ws3.cell(row=lr, column=3).number_format = '0.00'
    ws3.cell(row=lr, column=3).alignment = Alignment(horizontal="center")

    wb.save(path)   # Guardar el archivo Excel en disco
    return tots_ton, tots_km, tots_viaj, used_cams


# ──────────────────────────────────────────────────────────────
#  SECCIÓN 9 — MAPA VISUAL (matplotlib)
# ──────────────────────────────────────────────────────────────

def make_map(camiones, path):
    """
    Genera un mapa PNG de las rutas sobre un fondo oscuro.
    Izquierda: mapa geográfico con las rutas.
    Derecha: tabla de resumen por camión.
    """
    fig = plt.figure(figsize=(24, 13), facecolor="#0D1B2A")   # Figura 24×13 pulgadas

    # Crear dos ejes: el mapa ocupa 63% del ancho, la tabla el 33% restante.
    # add_axes([left, bottom, width, height]) en fracción de la figura (0 a 1).
    ax  = fig.add_axes([0.01, 0.07, 0.63, 0.87])   # Mapa (izquierda)
    ax2 = fig.add_axes([0.66, 0.07, 0.33, 0.87])   # Tabla resumen (derecha)
    for a in [ax, ax2]: a.set_facecolor("#0D1B2A")  # Fondo azul oscuro
    ax2.axis("off")   # Ocultar ejes de la tabla (no se necesitan)

    dep_lat = DEPOSITO["lat"]; dep_lon = DEPOSITO["lon"]
    legend_h = []   # Acumula parches de colores para la leyenda
    resumen  = []   # Acumula datos para la tabla derecha

    for ci, cam in enumerate(camiones):
        viajes = [v for v in cam if v["nodos"]]
        if not viajes:
            continue   # Camión sin rutas → no dibujar nada

        color = PALETTE[ci % len(PALETTE)]
        ton_t = round(sum(v["carga"] for v in viajes), 2)
        km_t  = round(sum(dist_viaje(v["nodos"]) for v in viajes), 2)
        resumen.append((ci + 1, len(viajes), ton_t, km_t, color))

        for vi, viaje in enumerate(viajes):
            ns   = viaje["nodos"]
            # Construir la secuencia de lat/lon para dibujar la ruta:
            # depósito → cada colonia → depósito
            lats = [dep_lat] + [COORDS[n-1][0] for n in ns] + [dep_lat]
            lons = [dep_lon] + [COORDS[n-1][1] for n in ns] + [dep_lon]

            # El primer viaje es más grueso y opaco; los siguientes más tenues
            lw    = 2.0 - vi * 0.5    # Grosor de línea: 2.0, 1.5, 1.0
            alpha = 0.9 - vi * 0.12   # Transparencia: 0.90, 0.78, 0.66

            # Dibujar la línea de la ruta
            ax.plot(lons, lats, color=color, linewidth=lw, alpha=alpha,
                    solid_capstyle="round", zorder=3)
            # zorder controla qué se dibuja encima; mayor zorder = al frente

            # Añadir flechas de dirección en el punto medio de cada tramo
            for k in range(len(lons) - 1):
                mx = (lons[k] + lons[k+1]) / 2   # Longitud del punto medio
                my = (lats[k] + lats[k+1]) / 2   # Latitud del punto medio
                dx = lons[k+1] - lons[k]          # Dirección en longitud
                dy = lats[k+1] - lats[k]          # Dirección en latitud
                # annotate dibuja una flecha de xytext → xy
                ax.annotate("",
                            xy=(mx + dx*0.015, my + dy*0.015),
                            xytext=(mx - dx*0.015, my - dy*0.015),
                            arrowprops=dict(arrowstyle="->", color=color,
                                           lw=0.7, mutation_scale=7),
                            zorder=5)

        # Marcar con punto las colonias asignadas a este camión
        # set() evita marcar dos veces si la misma colonia aparece en varios viajes
        for n in set(n2 for v in viajes for n2 in v["nodos"]):
            ax.scatter(COORDS[n-1][1], COORDS[n-1][0],
                       s=40, color=color, zorder=6,
                       edgecolors="white", linewidths=0.6)
            # Etiqueta con el nombre de la colonia, encima del punto
            ax.text(COORDS[n-1][1], COORDS[n-1][0] + 0.0015,
                    COLONIAS[n-1], fontsize=4.5, color="#DDDDDD",
                    ha="center", va="bottom", zorder=7,
                    fontfamily="DejaVu Sans Mono")

        # Agregar este camión a la leyenda
        lbl = f"C{ci+1:02d} · {len(viajes)}v · {ton_t}t"
        legend_h.append(mpatches.Patch(color=color, label=lbl))

    # Marcar el depósito con una estrella dorada grande
    ax.scatter(dep_lon, dep_lat, s=280, color="#FFD700", marker="*",
               zorder=10, edgecolors="white", linewidths=1.2)
    ax.text(dep_lon, dep_lat + 0.0035, "Paso Texca\n(Depósito)",
            fontsize=7.5, color="#FFD700", ha="center",
            fontweight="bold", zorder=11)

    # Títulos y estilos del mapa
    ax.set_title("Rutas de Recolección de Basura — Puerto de Acapulco",
                 color="white", fontsize=12, fontweight="bold", pad=10)
    ax.set_xlabel("Longitud", color="#888888", fontsize=8)
    ax.set_ylabel("Latitud",  color="#888888", fontsize=8)
    ax.tick_params(colors="#888888", labelsize=7)
    for sp in ax.spines.values(): sp.set_edgecolor("#2A2A4A")

    # Leyenda en la esquina inferior izquierda, 2 columnas
    ax.legend(handles=legend_h, loc="lower left", fontsize=5.2, ncol=2,
              framealpha=0.35, facecolor="#1A1A2E", edgecolor="#444",
              labelcolor="white", title="Cam · viajes · ton", title_fontsize=5.5)

    # ── Tabla de resumen en el panel derecho ──────────────────
    ax2.set_title("Resumen por camión", color="white",
                  fontsize=11, fontweight="bold", pad=8)

    col_labels = ["Cam.", "Viajes", "Ton.", "Km est."]
    tdata = [[f"C{d[0]:02d}", d[1], f"{d[2]} t", f"{d[3]} km"]
             for d in resumen]

    # ax2.table() dibuja una tabla dentro del eje (sin usar Excel)
    tbl = ax2.table(cellText=tdata, colLabels=col_labels,
                    loc="center", cellLoc="center")
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(7.5)
    tbl.scale(1.0, 1.28)   # (ancho_factor, alto_factor)

    # Colorear encabezados de la tabla
    for j in range(4):
        tbl[(0, j)].set_facecolor("#1F3A5F")
        tbl[(0, j)].get_text().set_color("white")
        tbl[(0, j)].get_text().set_fontweight("bold")

    # Colorear filas de datos: col 0 con el color del camión, resto alternado
    for i, d in enumerate(resumen):
        c  = d[4]                                    # Color del camión
        bg = "#1A2A3A" if i % 2 == 0 else "#0F1E2E" # Filas alternadas
        tbl[(i+1, 0)].set_facecolor(c)
        tbl[(i+1, 0)].get_text().set_color("white")
        tbl[(i+1, 0)].get_text().set_fontweight("bold")
        for j in range(1, 4):
            tbl[(i+1, j)].set_facecolor(bg)
            tbl[(i+1, j)].get_text().set_color("#DDDDDD")

    # Texto de totales en el pie de la figura
    tvj  = sum(d[1] for d in resumen)
    tton = round(sum(d[2] for d in resumen), 1)
    tkm  = round(sum(d[3] for d in resumen), 1)
    fig.text(0.5, 0.012,
             f"Total: {len(resumen)} camiones activos   ·   {tvj} viajes   ·   "
             f"{tton} t recolectadas   ·   {tkm} km totales estimados",
             ha="center", color="#AAAAAA", fontsize=9,
             bbox=dict(boxstyle="round,pad=0.4",
                       facecolor="#1A1A2E", edgecolor="#444"))

    # Guardar imagen con 165 ppp de resolución (calidad alta para impresión)
    plt.savefig(path, dpi=165, bbox_inches="tight", facecolor="#0D1B2A")
    plt.close()   # Liberar memoria
    print(f"  Mapa → {path}")


# ──────────────────────────────────────────────────────────────
#  SECCIÓN 10 — PUNTO DE ENTRADA (MAIN)
# ──────────────────────────────────────────────────────────────
#
#  `if __name__ == "__main__"` asegura que este bloque solo se ejecute
#  cuando se corre el script directamente (no cuando se importa como módulo).

if __name__ == "__main__":
    # Mensaje inicial con los parámetros del problema
    print("=" * 60)
    print("  CVRP · Algoritmo Genético · Acapulco v3")
    print("=" * 60)
    print(f"  Colonias:{NUM_COL}  Camiones:{NUM_CAM}  Cap:{CAP}t  "
          f"Máx viajes:{MAX_VIAJES}  Total:{TOTAL_BAS}t")
    print(f"  Demanda/colonia: {min(DEMANDA):.2f}–{max(DEMANDA):.2f} t\n")

    # ── PASO 1: Ejecutar el Algoritmo Genético ────────────────
    # Retorna el mejor cromosoma encontrado tras 500 generaciones
    mejor = ga(tam, gens, elite, seed)

    # ── PASO 2: Convertir el cromosoma en rutas reales ────────
    cams = decodificar(mejor)
    # cams[i] = lista de viajes del camión i
    # cams[i][j] = {"nodos": [lista de colonias], "carga": float}

    # Rutas de salida
    OUT_X = "/mnt/user-data/outputs/ruteo_acapulco.xlsx"
    OUT_P = "/mnt/user-data/outputs/mapa_rutas_acapulco.png"

    # ── PASO 3: Generar Excel ─────────────────────────────────
    print("\n  Generando Excel...")
    tt, tk, tv, uc = make_excel(cams, OUT_X)
    # tt = lista de toneladas por camión
    # tk = lista de km por camión
    # tv = lista de viajes por camión
    # uc = número de camiones activos

    # ── PASO 4: Generar mapa PNG ──────────────────────────────
    print("  Generando mapa...")
    make_map(cams, OUT_P)

    # ── PASO 5: Imprimir resumen en consola ───────────────────
    print("\n" + "=" * 60)
    print("  RESUMEN FINAL")
    print(f"  {'CAM':<8}{'VIAJES':>7}{'TON':>9}{'KM':>9}")
    print("  " + "-" * 34)
    for i in range(NUM_CAM):
        if tv[i] == 0:
            continue   # Omitir camiones sin actividad
        print(f"  C{i+1:02d}    {tv[i]:>6}  {tt[i]:>8.2f}  {tk[i]:>8.2f}")
    print("  " + "-" * 34)
    print(f"  TOTAL   {sum(tv):>6}  {sum(tt):>8.2f}  {sum(tk):>8.2f}")
    print(f"  Camiones usados: {uc}/{NUM_CAM}")
    print("=" * 60)
    print(f"\n  ✓ Excel → {OUT_X}")
    print(f"  ✓ Mapa  → {OUT_P}\n")
